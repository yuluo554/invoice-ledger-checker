"""技术报告 docx 生成器：docs/技术报告.md -> docs/技术报告.docx。

纪律（plan/05 §4 交付对标）：md 为唯一定稿源，docx 由本脚本程序化生成、
禁止手改产物；**先定稿 md 再生成**。用法：

    py -m pip install -e ".[report]"   # python-docx（extras: report）
    py docs/make_report_docx.py        # 仓库根执行；产物 docs/技术报告.docx

仅支持本项目报告用到的 markdown 子集：#/##/### 标题、段落（含 **加粗** 与
`行内码`）、- 列表、> 引用、| 表格 |、``` 围栏代码块。产物不入仓。
"""

import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(ROOT, "docs", "技术报告.md")
DOCX_PATH = os.path.join(ROOT, "docs", "技术报告.docx")


def _require_docx():
    try:
        import docx  # noqa: F401
    except ImportError as exc:
        raise SystemExit(
            "缺少 python-docx，请先安装：py -m pip install -e \".[report]\""
        ) from exc


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_CODE_RE = re.compile(r"`([^`]+)`")


def _add_runs(paragraph, text):
    """行内 **加粗** / `代码` 渲染为 run 序列。"""
    tokens = []
    pos = 0
    pattern = re.compile(r"\*\*(.+?)\*\*|`([^`]+)`")
    for match in pattern.finditer(text):
        if match.start() > pos:
            tokens.append(("plain", text[pos:match.start()]))
        if match.group(1) is not None:
            tokens.append(("bold", match.group(1)))
        else:
            tokens.append(("code", match.group(2)))
        pos = match.end()
    if pos < len(text):
        tokens.append(("plain", text[pos:]))
    for kind, chunk in tokens:
        run = paragraph.add_run(chunk)
        if kind == "bold":
            run.bold = True
        elif kind == "code":
            run.font.name = "Consolas"


def _add_table(doc, rows):
    """markdown 表格 -> docx 表格（首行表头加粗）。"""
    from docx.enum.table import WD_TABLE_ALIGNMENT

    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    if not cells:
        return
    table = doc.add_table(rows=len(cells), cols=len(cells[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(cells):
        for c, value in enumerate(row):
            if c >= len(table.rows[r].cells):
                continue
            cell = table.rows[r].cells[c]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            _add_runs(paragraph, value)
            if r == 0:
                for run in paragraph.runs:
                    run.bold = True


def convert(md_path, docx_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    with open(md_path, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # 跳过收尾 ```
            for code_line in code_lines:
                p = doc.add_paragraph()
                run = p.add_run(code_line if code_line else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            _add_table(doc, rows)
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            title = doc.add_heading(stripped[2:], level=1)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            _add_runs(p, stripped[2:])
            for run in p.runs:
                run.italic = True
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, re.sub(r"^\d+\.\s", "", stripped))
        else:
            p = doc.add_paragraph()
            _add_runs(p, stripped)
        i += 1

    doc.save(docx_path)
    return docx_path


def main():
    _require_docx()
    if not os.path.isfile(MD_PATH):
        raise SystemExit("定稿源不存在: %s" % MD_PATH)
    out = convert(MD_PATH, DOCX_PATH)
    print("docx 已生成: %s（源: %s；产物禁止手改，重跑本脚本刷新）" % (out, MD_PATH))
    return 0


if __name__ == "__main__":
    sys.exit(main())
