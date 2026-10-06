"""Excel 台账导出（openpyxl，M5 交付，FR-20）。

- 双 sheet：发票台账（invoices 全量字段）+ 异常清单（findings 三级记录）；
- 条件格式标红：异常清单按级别整行着色（error=红 / suspicious=黄 /
  review=蓝），发票台账对带字段标记（field_flags 非空）的行浅黄提示；
- openpyxl 经 extras: export 惰性导入（utils.import_optional 纪律）；
- 金额以 Decimal 文本入库，导出转数值单元格（显示层，比较仍以应用层为准）。
docx 摘要报告为 P2（extras: report）。
"""

import os
from typing import Any, Dict, List

from ..storage import Ledger
from ..utils import import_optional

# 异常清单 sheet 条件格式配色（Excel 标准 bad/neutral/info 变体）
_LEVEL_STYLES = [
    # (级别显示名, 整行填充色, 字体色)
    ("确认异常", "FFC7CE", "9C0006"),   # 红
    ("疑似", "FFEB9C", "9C6500"),       # 黄
    ("待人工确认", "BDD7EE", "1F4E79"),  # 蓝
]

# 级别码 -> Excel 显示名（与 _LEVEL_STYLES 顺序一致）
_LEVEL_LABELS = {
    "error": "确认异常",
    "suspicious": "疑似",
    "review": "待人工确认",
}
# 规则中文名（Excel 呈现层映射；规则语义权威出处 plan/04 §3）
_RULE_NAMES = {
    "R-DUP-01": "精确重复",
    "R-DUP-02": "联合键重复",
    "R-DUP-03": "模糊重复",
    "R-SEQ-01": "连号拆分簇",
    "R-ARITH-01": "价税合计复核",
    "R-ARITH-02": "大写金额一致",
    "R-TIME-01": "开票晚于导入",
    "R-TIME-02": "跨期异常",
}

INVOICE_HEADERS = [
    "发票号码", "发票类型", "开票日期", "购买方", "购买方税号",
    "销售方", "销售方税号", "金额（不含税）", "税额", "价税合计",
    "价税合计大写", "备注", "明细数", "批次", "置信度", "字段标记",
]
FINDING_HEADERS = [
    "finding_id", "规则ID", "规则名称", "级别", "关联发票号码", "说明",
    "批次", "状态", "检出时间",
]


def _to_number(text: str):
    """Decimal 文本 -> float（空串/非法原样返回，不造默认值）。"""
    if not text:
        return ""
    try:
        from decimal import Decimal

        return float(Decimal(text))
    except Exception:
        return text


def _style_header(sheet) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill("solid", fgColor="4472C4")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def _fit_columns(sheet, widths: Dict[int, int]) -> None:
    for idx, width in widths.items():
        sheet.column_dimensions[
            sheet.cell(row=1, column=idx).column_letter].width = width


def _add_level_conditional_formatting(sheet, last_row: int) -> None:
    """异常清单条件格式：级别列命中即整行着色（DoD：条件格式标红用例断言）。"""
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill

    last_col = len(FINDING_HEADERS)
    end = "%s%d" % (sheet.cell(row=1, column=last_col).column_letter, last_row)
    col = sheet.cell(row=1, column=FINDING_HEADERS.index("级别") + 1).column_letter
    for label, fill_color, font_color in _LEVEL_STYLES:
        sheet.conditional_formatting.add(
            "A2:%s" % end,
            FormulaRule(
                formula=['$%s2="%s"' % (col, label)],
                fill=PatternFill("solid", bgColor=fill_color),
                font=Font(color=font_color),
            ),
        )


def _add_flag_conditional_formatting(sheet, last_row: int) -> None:
    """发票台账条件格式：字段标记列非空 -> 整行浅黄（低置信度人工留意）。"""
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill

    last_col = len(INVOICE_HEADERS)
    end = "%s%d" % (sheet.cell(row=1, column=last_col).column_letter, last_row)
    col = sheet.cell(row=1, column=INVOICE_HEADERS.index("字段标记") + 1).column_letter
    sheet.conditional_formatting.add(
        "A2:%s" % end,
        FormulaRule(
            formula=['$%s2<>""' % col],
            fill=PatternFill("solid", bgColor="FFF2CC"),
            font=Font(color="7F6000"),
        ),
    )


def _fill_invoice_sheet(sheet, cards) -> None:
    sheet.append(INVOICE_HEADERS)
    for card in cards:
        sheet.append([
            card.invoice_number,
            card.invoice_type,
            card.issue_date,
            card.buyer_name,
            card.buyer_tax_id,
            card.seller_name,
            card.seller_tax_id,
            _to_number(card.amount),
            _to_number(card.tax_amount),
            _to_number(card.total_with_tax),
            card.total_with_tax_cn,
            card.remark,
            len(card.items),
            card.batch_id,
            card.confidence,
            "、".join(card.field_flags),
        ])
    for row in sheet.iter_rows(min_row=2):
        for idx in (8, 9, 10):  # 金额/税额/价税合计
            row[idx - 1].number_format = "#,##0.00"
    _style_header(sheet)
    _fit_columns(sheet, {1: 24, 3: 12, 4: 26, 5: 20, 6: 26, 7: 20, 11: 16, 16: 20})
    _add_flag_conditional_formatting(sheet, max(sheet.max_row, 2))


def _fill_finding_sheet(sheet, findings) -> None:
    sheet.append(FINDING_HEADERS)
    for finding in findings:
        sheet.append([
            finding["finding_id"],
            finding["rule_id"],
            _RULE_NAMES.get(finding["rule_id"], finding["rule_id"]),
            _LEVEL_LABELS.get(finding["level"], finding["level"]),
            "、".join(finding["invoice_numbers"]),
            finding["message"],
            finding.get("batch_id") or "",
            finding.get("status", "open"),
            finding.get("created_at", ""),
        ])
    _style_header(sheet)
    _fit_columns(sheet, {1: 44, 3: 14, 4: 12, 5: 46, 6: 40, 9: 20})
    _add_level_conditional_formatting(sheet, max(sheet.max_row, 2))


def export_ledger(db_path: str, out_path: str) -> Dict[str, Any]:
    """SQLite 台账 -> Excel 工作簿（发票台账 + 异常清单双 sheet）。

    返回摘要 dict（out_path/invoice_count/finding_count/levels 计数）。
    db 文件不存在抛 ValueError（CLI 转参数错误退出码 2）；缺 openpyxl 抛
    OptionalDependencyError（惰性导入，带安装提示）。
    """
    if not os.path.isfile(db_path):
        raise ValueError("台账数据库不存在: %s" % db_path)
    import_optional("openpyxl", "export")
    from openpyxl import Workbook

    ledger = Ledger(db_path)
    try:
        cards = list(ledger.iter_cards())
        findings = ledger.list_findings()
    finally:
        ledger.close()

    wb = Workbook()
    invoice_sheet = wb.active
    invoice_sheet.title = "发票台账"
    _fill_invoice_sheet(invoice_sheet, cards)
    finding_sheet = wb.create_sheet("异常清单")
    _fill_finding_sheet(finding_sheet, findings)
    wb.save(out_path)

    levels: Dict[str, int] = {}
    for finding in findings:
        levels[finding["level"]] = levels.get(finding["level"], 0) + 1
    return {
        "out_path": out_path,
        "invoice_count": len(cards),
        "finding_count": len(findings),
        "levels": levels,
        "sheets": ["发票台账", "异常清单"],
    }


def export_findings_xlsx(findings: List[Dict[str, Any]], out_path: str,
                         title: str = "异常清单") -> Dict[str, Any]:
    """仅异常清单单 sheet 导出（GUI 导出对话框复用同一渲染）。"""
    import_optional("openpyxl", "export")
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet.title = title[:31]  # Excel sheet 名上限 31 字符
    _fill_finding_sheet(sheet, findings)
    wb.save(out_path)
    return {"out_path": out_path, "finding_count": len(findings),
            "sheets": [sheet.title]}
